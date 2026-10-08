"""Controls-grid responsibility projection from workbook Columns L/M."""

from openpyxl import Workbook as OpenPyxlWorkbook

from cybersecurity_assessor.excel.ccis_reader import _INDEX_CACHE
from cybersecurity_assessor.routes.workbooks import workbook_col_l_status


def _write_responsibility_row(path, *, column_d: str, column_l: str, column_m: str | None):
    workbook = OpenPyxlWorkbook()
    sheet = workbook.active
    sheet.title = "WORKING SHEET"
    sheet.cell(6, 1).value = "Required"
    sheet.cell(6, 2).value = "Control Acronym"
    sheet.cell(7, 1).value = "YES"
    sheet.cell(7, 2).value = "AC-4"
    sheet.cell(7, 4).value = column_d
    sheet.cell(7, 8).value = "CCI-001548"
    sheet.cell(7, 12).value = column_l
    sheet.cell(7, 13).value = column_m
    workbook.save(path)
    _INDEX_CACHE.clear()


def test_inherited_rollup_returns_column_m_source(session, controls_catalog):
    _write_responsibility_row(
        controls_catalog["path"],
        column_d="Inherited",
        column_l="Remote",
        column_m="DoW Enterprise",
    )

    result = workbook_col_l_status(controls_catalog["workbook"].id, session)

    assert result == [
        {
            "control_id": "ac-4",
            "outcome": "inherited",
            "value": "Remote",
            "source": "DoW Enterprise",
        }
    ]


def test_column_d_na_omits_responsibility_source(session, controls_catalog):
    _write_responsibility_row(
        controls_catalog["path"],
        column_d="Not Applicable",
        column_l="Remote",
        column_m="DoW Enterprise",
    )

    result = workbook_col_l_status(controls_catalog["workbook"].id, session)

    assert result[0]["outcome"] == "na"
    assert result[0]["source"] is None
