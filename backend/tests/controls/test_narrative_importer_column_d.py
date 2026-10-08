"""Column-D authority for operator-supplied narrative imports."""

from __future__ import annotations

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook as OpenPyxlWorkbook
from sqlmodel import select

from cybersecurity_assessor.db import get_session
from cybersecurity_assessor.excel.ccis_reader import _INDEX_CACHE
from cybersecurity_assessor.excel.narrative_importer import import_narratives
from cybersecurity_assessor.models import (
    Assessment,
    ComplianceStatus,
    VerdictSource,
)
from cybersecurity_assessor.server import create_app


def _write_row(
    path,
    *,
    implementation_status: str,
    compliance_status: str | None,
    narrative: str | None,
) -> None:
    workbook = OpenPyxlWorkbook()
    sheet = workbook.active
    sheet.title = "WORKING SHEET"
    sheet.cell(6, 1).value = "Required"
    sheet.cell(6, 2).value = "Control Acronym"
    sheet.cell(7, 1).value = "YES"
    sheet.cell(7, 2).value = "AC-2"
    sheet.cell(7, 4).value = implementation_status
    sheet.cell(7, 8).value = "CCI-000015"
    sheet.cell(7, 14).value = compliance_status
    sheet.cell(7, 16).value = "Import Assessor"
    sheet.cell(7, 17).value = narrative
    workbook.save(path)


def test_source_column_d_overrides_imported_compliant_status(
    session, controls_catalog, tmp_path
):
    source_path = controls_catalog["path"]
    _write_row(
        source_path,
        implementation_status="Not Applicable",
        compliance_status=None,
        narrative="Not required for GMI CUI DIT Environment.",
    )
    import_path = tmp_path / "operator-results.xlsx"
    _write_row(
        import_path,
        implementation_status="Implemented",
        compliance_status="Compliant",
        narrative="Operator import incorrectly labels this Compliant.",
    )
    _INDEX_CACHE.clear()

    result = import_narratives(
        session,
        controls_catalog["workbook"].id,
        import_path,
    )

    assessment = session.exec(select(Assessment)).one()
    assert result.imported == 1
    assert result.overridden_by_column_d == 1
    assert assessment.status is ComplianceStatus.NOT_APPLICABLE
    assert assessment.inheritance_rule == "8b"
    assert assessment.verdict_source is VerdictSource.RULE_8B


def test_non_na_source_column_d_preserves_imported_status(
    session, controls_catalog, tmp_path
):
    source_path = controls_catalog["path"]
    _write_row(
        source_path,
        implementation_status="Planned",
        compliance_status=None,
        narrative=None,
    )
    import_path = tmp_path / "operator-results.xlsx"
    _write_row(
        import_path,
        implementation_status="Planned",
        compliance_status="Compliant",
        narrative="Verified via the approved account management plan.",
    )
    _INDEX_CACHE.clear()

    result = import_narratives(
        session,
        controls_catalog["workbook"].id,
        import_path,
    )

    assessment = session.exec(select(Assessment)).one()
    assert result.overridden_by_column_d == 0
    assert assessment.status is ComplianceStatus.COMPLIANT
    assert assessment.inheritance_rule is None
    assert assessment.verdict_source is VerdictSource.IMPORTED


def test_non_na_source_column_d_rejects_imported_na(
    session, controls_catalog, tmp_path
):
    _write_row(
        controls_catalog["path"],
        implementation_status="Planned",
        compliance_status=None,
        narrative=None,
    )
    import_path = tmp_path / "operator-na.xlsx"
    _write_row(
        import_path,
        implementation_status="Not Applicable",
        compliance_status="Not Applicable",
        narrative="Operator attempted to scope the CCI out.",
    )
    _INDEX_CACHE.clear()

    with pytest.raises(ValueError, match="conflicts with source workbook Column D"):
        import_narratives(
            session,
            controls_catalog["workbook"].id,
            import_path,
        )

    assert session.exec(select(Assessment)).all() == []


def test_column_d_na_import_does_not_require_narrative(
    session, controls_catalog, tmp_path
):
    _write_row(
        controls_catalog["path"],
        implementation_status="Not Applicable",
        compliance_status=None,
        narrative=None,
    )
    import_path = tmp_path / "blank-result.xlsx"
    _write_row(
        import_path,
        implementation_status="Implemented",
        compliance_status="Compliant",
        narrative=None,
    )
    _INDEX_CACHE.clear()

    result = import_narratives(
        session,
        controls_catalog["workbook"].id,
        import_path,
    )

    assessment = session.exec(select(Assessment)).one()
    assert result.imported == 1
    assert result.skipped_no_narrative == []
    assert assessment.status is ComplianceStatus.NOT_APPLICABLE
    assert "Column D" in assessment.narrative_q


def test_non_na_import_without_narrative_remains_skipped(
    session, controls_catalog, tmp_path
):
    _write_row(
        controls_catalog["path"],
        implementation_status="Planned",
        compliance_status=None,
        narrative="Old program-workbook narrative must not be borrowed.",
    )
    import_path = tmp_path / "blank-non-na-result.xlsx"
    _write_row(
        import_path,
        implementation_status="Planned",
        compliance_status="Compliant",
        narrative=None,
    )
    _INDEX_CACHE.clear()

    result = import_narratives(
        session,
        controls_catalog["workbook"].id,
        import_path,
    )

    assert result.imported == 0
    assert result.skipped_no_narrative == ["CCI-000015"]
    assert session.exec(select(Assessment)).all() == []


def test_manual_force_cannot_override_source_column_d_na(
    session, controls_catalog
):
    _write_row(
        controls_catalog["path"],
        implementation_status="Not Applicable",
        compliance_status=None,
        narrative="Not required for GMI CUI DIT Environment.",
    )
    _INDEX_CACHE.clear()

    app = create_app()

    def _override_session():
        yield session

    app.dependency_overrides[get_session] = _override_session
    objective = controls_catalog["objectives"]["CCI-000015"]
    response = TestClient(app).post(
        "/api/controls/assessments?force=true",
        json={
            "workbook_id": controls_catalog["workbook"].id,
            "objective_id": objective.id,
            "status": "Compliant",
            "tester": "Noah Jaskolski",
            "narrative_q": "Verified via an implementation artifact.",
            "narrative_class": "compliance-affirming",
        },
    )

    assert response.status_code == 409
    assert "Column D" in response.json()["detail"]
    assert session.exec(select(Assessment)).all() == []


def test_manual_force_cannot_create_na_when_source_column_d_is_planned(
    session, controls_catalog
):
    _write_row(
        controls_catalog["path"],
        implementation_status="Planned",
        compliance_status=None,
        narrative=None,
    )
    _INDEX_CACHE.clear()

    app = create_app()

    def _override_session():
        yield session

    app.dependency_overrides[get_session] = _override_session
    objective = controls_catalog["objectives"]["CCI-000015"]
    response = TestClient(app).post(
        "/api/controls/assessments?force=true",
        json={
            "workbook_id": controls_catalog["workbook"].id,
            "objective_id": objective.id,
            "status": "Not Applicable",
            "tester": "Noah Jaskolski",
            "narrative_q": "Not applicable because the operator said so.",
            "narrative_class": "NA-justifying",
        },
    )

    assert response.status_code == 409
    assert "Column D" in response.json()["detail"]
    assert session.exec(select(Assessment)).all() == []


def test_control_detail_fails_closed_when_source_workbook_is_missing(
    session, controls_catalog
):
    controls_catalog["path"].unlink()
    _INDEX_CACHE.clear()

    app = create_app()

    def _override_session():
        yield session

    app.dependency_overrides[get_session] = _override_session
    control = controls_catalog["controls"]["AC-2"]
    response = TestClient(app).get(
        f"/api/controls/{control.id}",
        params={"workbook_id": controls_catalog["workbook"].id},
    )

    assert response.status_code == 410
    assert "Source workbook not found" in response.json()["detail"]
