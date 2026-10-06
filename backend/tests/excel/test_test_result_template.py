"""Regression coverage for eMASS Test Result Import templates."""

from __future__ import annotations

from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from cybersecurity_assessor.excel.ccis_reader import (
    _INDEX_CACHE,
    read_workbook_index,
)
from cybersecurity_assessor.excel.ccis_validator import validate_workbook
from cybersecurity_assessor.excel.ccis_writer import write_single


def _build_template(path, *, row_count: int = 520):
    wb = Workbook()
    ws = wb.active
    ws.title = "Template"

    headers = {
        1: "Control Acronym",
        2: "Control Set",
        3: "Control Information",
        4: "Control Implementation Status",
        5: "Security Control Designation",
        6: "Control Implementation Narrative",
        7: "AP Acronym",
        8: "CCI",
        9: "CCI Definition",
        10: "Implementation Guidance",
        11: "Assessment Procedures",
        12: "Inherited",
        13: "Remote Inheritance Instance",
        14: "Compliance Status",
        15: "Date Tested",
        16: "Tested By",
        17: "Test Results",
        18: "Compliance Status",
        19: "Date Tested",
        20: "Tested By",
        21: "Test Results",
    }
    for col, value in headers.items():
        ws.cell(row=6, column=col, value=value)

    for offset in range(row_count):
        row = 7 + offset
        ws.cell(row=row, column=1, value="AC-1")
        ws.cell(row=row, column=2, value="NIST SP 800-53 Revision 4")
        ws.cell(row=row, column=4, value="Implemented")
        ws.cell(row=row, column=7, value=f"AC-1.{offset + 1}")
        ws.cell(row=row, column=8, value=offset + 1)
        ws.cell(row=row, column=9, value=f"Objective {offset + 1}")
        ws.cell(row=row, column=12, value="Local")

    ws.cell(row=7, column=14, value="Compliant")
    ws.cell(row=7, column=15, value="10/05/2026 1:02:03 PM")
    ws.cell(row=7, column=16, value="Assessor")
    ws.cell(row=7, column=17, value="Verified via test artifact.")
    wb.create_sheet("Example")
    wb.save(path)
    return path


def test_reader_accepts_large_test_result_template(tmp_path):
    _INDEX_CACHE.clear()
    path = _build_template(tmp_path / "TRExport_Template.xlsx")

    index = read_workbook_index(path)

    assert index.sheet_name == "Template"
    assert len(index.rows) == 520
    assert index.rows[0].required is True
    assert index.rows[0].control_id == "AC-1"
    assert index.rows[0].cci_id == "CCI-000001"
    assert index.rows[0].date_tested == datetime(2026, 10, 5, 13, 2, 3)
    assert index.rows[-1].excel_row == 526


def test_validator_accepts_test_result_template_headers(tmp_path):
    path = _build_template(tmp_path / "TRExport_Template.xlsx", row_count=2)

    report = validate_workbook(path)

    assert report.sheet_name == "Template"
    assert report.valid is True
    assert report.errors == []


def test_writer_updates_template_assessment_columns(tmp_path):
    _INDEX_CACHE.clear()
    path = _build_template(tmp_path / "TRExport_Template.xlsx", row_count=2)

    result = write_single(
        path,
        excel_row=8,
        status="Non-Compliant",
        date_tested=datetime(2026, 10, 5),
        tester="Noah Jaskolski",
        results="No implementation evidence was located; POA&M required.",
    )

    assert result["sheet"] == "Template"
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb["Template"]
        assert ws["A8"].value == "AC-1"
        assert ws["N8"].value == "Non-Compliant"
        assert ws["O8"].value == "2026-10-05"
        assert ws["P8"].value == "Noah Jaskolski"
        assert ws["Q8"].value == (
            "No implementation evidence was located; POA&M required."
        )
    finally:
        wb.close()


def test_reader_keeps_rows_after_internal_blank_run(tmp_path):
    _INDEX_CACHE.clear()
    path = _build_template(tmp_path / "gapped.xlsx", row_count=1)
    wb = load_workbook(path)
    ws = wb["Template"]
    ws.cell(row=14, column=1, value="AC-2")
    ws.cell(row=14, column=7, value="AC-2.1")
    ws.cell(row=14, column=8, value=15)
    wb.save(path)

    index = read_workbook_index(path)
    assert [row.excel_row for row in index.rows] == [7, 14]
    assert index.rows[-1].cci_id == "CCI-000015"


def test_reader_rejects_sheet_beyond_safety_limit(tmp_path, monkeypatch):
    _INDEX_CACHE.clear()
    monkeypatch.setattr(
        "cybersecurity_assessor.excel.ccis_reader._MAX_DATA_ROW", 20
    )
    path = _build_template(tmp_path / "oversized.xlsx", row_count=1)
    wb = load_workbook(path)
    ws = wb["Template"]
    ws.cell(row=21, column=1, value="AC-2")
    ws.cell(row=21, column=8, value=15)
    wb.save(path)

    with pytest.raises(ValueError, match="Refusing to silently truncate"):
        read_workbook_index(path)
