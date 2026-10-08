"""Column-D authority at single and bulk workbook-write boundaries."""

# ruff: noqa: F811

from __future__ import annotations

from openpyxl import Workbook as OpenPyxlWorkbook

from cybersecurity_assessor.excel.ccis_reader import _INDEX_CACHE
from cybersecurity_assessor.models import ComplianceStatus
from cybersecurity_assessor.routes import controls as controls_route
from cybersecurity_assessor.routes.controls import (
    ApplyBatchBody,
    ApplyToWorkbookBody,
    apply_assessment_to_workbook,
    apply_assessments_batch_to_workbook,
)
from tests.controls.conftest import assess, controls_catalog, session  # noqa: F401


def _write_column_d_na(path) -> None:
    workbook = OpenPyxlWorkbook()
    sheet = workbook.active
    sheet.title = "WORKING SHEET"
    sheet.cell(6, 1).value = "Required"
    sheet.cell(6, 2).value = "Control Acronym"
    sheet.cell(7, 1).value = "YES"
    sheet.cell(7, 2).value = "AC-2"
    sheet.cell(7, 4).value = "Not Applicable"
    sheet.cell(7, 8).value = "CCI-000015"
    sheet.cell(7, 21).value = "Not required for GMI CUI DIT Environment."
    workbook.save(path)


def test_single_apply_writes_column_d_na_over_stale_compliant(
    session, controls_catalog, assess, tmp_path, monkeypatch
):
    _write_column_d_na(controls_catalog["path"])
    _INDEX_CACHE.clear()
    assessment = assess(
        controls_catalog["workbook"].id,
        controls_catalog["objectives"]["CCI-000015"].id,
        ComplianceStatus.COMPLIANT,
        narrative="Stale Compliant result.",
    )
    captured = {}
    monkeypatch.setattr(
        controls_route,
        "get_or_create_working_copy",
        lambda wb, s: tmp_path / "working.xlsx",
    )
    monkeypatch.setattr(
        controls_route.ccis_writer,
        "write_single",
        lambda path, **kwargs: captured.update(kwargs) or {"rows_written": 1},
    )

    apply_assessment_to_workbook(
        ApplyToWorkbookBody(assessment_id=assessment.id),
        session,
    )

    assert captured["status"] is ComplianceStatus.NOT_APPLICABLE
    assert "Column D" in captured["results"]


def test_batch_apply_writes_column_d_na_over_stale_compliant(
    session, controls_catalog, assess, tmp_path, monkeypatch
):
    _write_column_d_na(controls_catalog["path"])
    _INDEX_CACHE.clear()
    assessment = assess(
        controls_catalog["workbook"].id,
        controls_catalog["objectives"]["CCI-000015"].id,
        ComplianceStatus.COMPLIANT,
        narrative="Stale Compliant result.",
    )
    captured = []
    monkeypatch.setattr(
        controls_route,
        "get_or_create_working_copy",
        lambda wb, s: tmp_path / "working.xlsx",
    )
    monkeypatch.setattr(
        controls_route.ccis_writer,
        "write_assessment",
        lambda path, writes, **kwargs: captured.extend(writes) or {
            "rows_written": len(captured),
            "skipped_needs_review": 0,
        },
    )

    apply_assessments_batch_to_workbook(
        ApplyBatchBody(
            workbook_id=controls_catalog["workbook"].id,
            assessment_ids=[assessment.id],
        ),
        session,
    )

    assert len(captured) == 1
    assert captured[0].status is ComplianceStatus.NOT_APPLICABLE
    assert "Column D" in captured[0].results
